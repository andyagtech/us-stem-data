#!/usr/bin/env python3
"""
Extract NSF SDR (Survey of Doctorate Recipients) 2023 citizenship and field
data into a single JSON for the STEM workforce dashboard page.

Output: data/nsf_workforce/workforce_dashboard.json
"""
import json
import warnings
import openpyxl
from pathlib import Path

warnings.filterwarnings("ignore")

ROOT = Path(__file__).resolve().parent.parent
SDR = ROOT / "data" / "nsf_workforce" / "sdr_2023_tables"
OUT = ROOT / "data" / "nsf_workforce" / "workforce_dashboard.json"


def load_table9_citizenship_by_field():
    """Table 9: U.S. residing employed doctoral S&E by field × citizenship status."""
    wb = openpyxl.load_workbook(SDR / "nsf25321-tab009.xlsx", data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    # Header layout (row indices 3-5):
    # Row 3: Field of study | All employed |   | U.S. citizen ... | Non-U.S. citizen ...
    # Row 4: blank | blank | blank | Total | blank | Native born | blank | Naturalized | blank | Total | blank | Permanent resident | blank | Temporary resident | blank
    # Row 5: numeric column headers

    # We want broad fields and a few key fine fields
    broad_fields = {
        "All fields": "All Fields",
        "Science": "Science (all)",
        "Engineering": "Engineering (all)",
        "Health": "Health Sciences",
        "Biological, agricultural, and environmental life sciences": "Life Sciences",
        "Computer and information sciences": "Computer & Information Sciences",
        "Computer science": "Computer Science",
        "Mathematics and statistics": "Mathematics & Statistics",
        "Physics": "Physics",
        "Chemistry, except biochemistry": "Chemistry",
        "Psychology": "Psychology",
        "Social sciences": "Social Sciences",
        "Aerospace, aeronautical, and astronautical engineering": "Aerospace Eng",
        "Chemical engineering": "Chemical Eng",
        "Civil engineering": "Civil Eng",
        "Electrical and computer engineering": "Electrical & Computer Eng",
        "Mechanical engineering": "Mechanical Eng",
        "Bioengineering and biomedical engineering": "Biomedical Eng",
    }
    out = []
    for r in rows:
        if not r or not r[0] or not isinstance(r[0], str):
            continue
        label = r[0].strip()
        if label not in broad_fields:
            continue
        try:
            total = int(r[1] or 0)
            us_total = int(r[3] or 0)
            native = int(r[5] or 0)
            naturalized = int(r[7] or 0)
            non_us_total = int(r[9] or 0)
            perm_res = int(r[11] or 0)
            temp_res = int(r[13] or 0)
        except (TypeError, ValueError):
            continue
        out.append({
            "field": broad_fields[label],
            "total": total,
            "native_born": native,
            "naturalized": naturalized,
            "permanent_resident": perm_res,
            "temporary_resident": temp_res,
            # Derived: foreign-born share = (naturalized + perm_res + temp_res) / total
            "foreign_born_pct": round(100 * (naturalized + perm_res + temp_res) / total, 1) if total else 0,
            "non_citizen_pct": round(100 * non_us_total / total, 1) if total else 0,
        })
    return out


def load_existing_findings():
    """The national topline, with every share RECOMPUTED from the counts.

    The input file carries typed `pct` fields and a typed `key_finding`
    sentence, and they disagreed with the counts beside them. The sentence said
    "22% naturalized citizens + 7% green card holders" where the file's own
    numbers give 20.2% and 7.9%, and it called 63.5% the native-born share "of
    PhD holders in the US" when 63.5% is the share of ALL respondents including
    the 8.6% living abroad -- the US-based share is 69.5%. /workforce.html
    quoted the sentence faithfully, so the page was wrong because the dataset
    was.

    So nothing typed survives. Counts are the input; every share and the
    sentence itself come out of this function.

    THREE UNIVERSES, which must never be mixed in one sentence:
      * all respondents          -- the survey sample, including those abroad
      * US-based respondents     -- the sample minus those abroad
      * the IPEDS domestic bucket -- native + naturalised + permanent resident,
        which is the group IPEDS reports as one number
    """
    fp = ROOT / "data" / "nsf_workforce" / "sdr_citizenship_findings.json"
    if not fp.exists():
        return None
    raw = json.loads(fp.read_text())
    cb = raw["citizenship_breakdown"]
    n = {k: int(v["count"]) for k, v in cb.items()}

    all_resp = sum(n.values())
    us_based = all_resp - n["outside_us"]
    ipeds_bucket = (n["us_citizen_native"] + n["us_citizen_naturalized"]
                    + n["permanent_resident"])
    foreign_in_bucket = n["us_citizen_naturalized"] + n["permanent_resident"]

    def pct(num, den):
        return round(num / den * 100, 1)

    for key, entry in cb.items():
        entry["pct"] = pct(n[key], all_resp)
        entry["pct_of_us_based"] = pct(n[key], us_based)

    raw["totals"] = {
        "all_respondents": all_resp,
        "us_based_respondents": us_based,
        "ipeds_domestic_bucket": ipeds_bucket,
    }
    raw["derived"] = {
        "native_born_pct_of_us_based": pct(n["us_citizen_native"], us_based),
        "naturalised_pct_of_us_based": pct(n["us_citizen_naturalized"], us_based),
        "permanent_resident_pct_of_us_based": pct(n["permanent_resident"], us_based),
        "temporary_visa_pct_of_us_based": pct(n["temporary_visa"], us_based),
        "foreign_born_pct_of_ipeds_bucket": pct(foreign_in_bucket, ipeds_bucket),
        "naturalised_pct_of_ipeds_bucket": pct(n["us_citizen_naturalized"], ipeds_bucket),
        "permanent_resident_pct_of_ipeds_bucket": pct(n["permanent_resident"], ipeds_bucket),
    }
    d = raw["derived"]
    raw["key_finding"] = (
        f"{d['foreign_born_pct_of_ipeds_bucket']}% of the group IPEDS reports as "
        f"US citizen or permanent resident is foreign-born: "
        f"{d['naturalised_pct_of_ipeds_bucket']}% naturalised citizens and "
        f"{d['permanent_resident_pct_of_ipeds_bucket']}% green card holders. "
        f"Of PhD holders based in the US, "
        f"{d['native_born_pct_of_us_based']}% are native-born Americans."
    )
    return raw


def load_nsb_stem_by_industry():
    """NSB STAL-12: STEM workforce by industry × place-of-birth + citizenship (2023)."""
    fp = ROOT / "data" / "nsf_workforce" / "nsb20261_stem_labor" / "nsb20261-tabstal-012.xlsx"
    if not fp.exists():
        return None
    wb = openpyxl.load_workbook(fp, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    out = []
    # Data rows start at row 6
    for r in rows[5:]:
        if not r or not r[0] or not isinstance(r[0], str):
            continue
        industry = r[0].strip()
        if industry in ("All industries", "Total", "Major industries"):
            continue
        try:
            all_workers = int(r[1] or 0)
            stem = int(r[2] or 0)
            us_born = int(r[3] or 0)
            foreign_born = int(r[4] or 0)
            naturalized = int(r[8] or 0)
            noncitizen = int(r[9] or 0)
        except (TypeError, ValueError):
            continue
        if stem == 0:
            continue
        out.append({
            "industry": industry,
            "all_workers": all_workers,
            "stem_workers": stem,
            "us_born": us_born,
            "foreign_born": foreign_born,
            "foreign_born_pct": round(100 * foreign_born / stem, 1) if stem else 0,
            "naturalized": naturalized,
            "noncitizen": noncitizen,
        })
    return out


def main():
    out = {
        "source": "NSF Survey of Doctorate Recipients (SDR) 2023",
        "url": "https://ncses.nsf.gov/surveys/doctorate-recipients",
        "national_topline": load_existing_findings(),
        "by_field": load_table9_citizenship_by_field(),
        "nsb_stem_by_industry": load_nsb_stem_by_industry(),
    }
    OUT.write_text(json.dumps(out, indent=2))
    print(f"Wrote {OUT}")
    print(f"  by_field rows: {len(out['by_field'])}")
    for f in out["by_field"]:
        print(f"    {f['field']:<32} {f['total']:>10,} ({f['foreign_born_pct']}% foreign-born)")
    if out["nsb_stem_by_industry"]:
        print(f"  NSB STEM industries: {len(out['nsb_stem_by_industry'])}")
        for i in out["nsb_stem_by_industry"][:5]:
            print(f"    {i['industry'][:40]:<42} {i['stem_workers']:>10,} STEM ({i['foreign_born_pct']}% foreign-born)")


if __name__ == "__main__":
    main()
