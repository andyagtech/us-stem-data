#!/usr/bin/env python3
"""What each occupation actually pays, from the national wage survey.

The H-1B filings say which prevailing-wage LEVEL an employer promised, I
through IV. That is a position in a distribution without saying what the
distribution is. This is the distribution: the 10th, 25th, 50th, 75th and
90th percentile annual wage for every occupation, nationally, from the
Occupational Employment and Wage Statistics survey.

On its own it establishes nothing about visas. It is the denominator that
makes the wage-level finding legible in dollars, and it is built separately
so it can be checked separately.

bls.gov refuses a request whose user agent contains a parenthetical, which is
what made this look unreachable for two sessions. A plain token works and no
key is needed for the bulk files.

Usage:
  python3 scripts/build_occupational_wages.py
"""
from __future__ import annotations

import io
import json
import re
import sys
import warnings
import zipfile
from pathlib import Path

warnings.filterwarnings("ignore")

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data" / "bls_oes"
OUT = ROOT / "site" / "data" / "occupational_wages.json"

# O_GROUP marks the level of aggregation; "detailed" is the six-digit SOC.
DETAIL = "detailed"
WAGE_COLS = {"p10": "A_PCT10", "p25": "A_PCT25", "median": "A_MEDIAN",
             "p75": "A_PCT75", "p90": "A_PCT90", "mean": "A_MEAN"}
# The occupations the H-1B programme actually runs through, so the output
# leads with them rather than with all 800.
FOCUS = {
    "15-1252": "Software developers",
    "15-1211": "Computer systems analysts",
    "15-1232": "Computer user support specialists",
    "15-1231": "Computer network support specialists",
    "15-1244": "Network and computer systems administrators",
    "15-2051": "Data scientists",
    "17-2061": "Computer hardware engineers",
    "17-2071": "Electrical engineers",
    "17-2141": "Mechanical engineers",
    "13-2011": "Accountants and auditors",
    "11-3021": "Computer and information systems managers",
    "25-1071": "Health specialties teachers, postsecondary",
    "29-1221": "Pediatricians, general",
}


def num(v):
    """OEWS writes '*' for suppressed and '#' for wages above $239,200."""
    s = str(v or "").strip().replace(",", "").replace("$", "")
    if s in ("", "*", "#", "**"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def read_year(path: Path) -> tuple[int, dict]:
    import openpyxl
    with zipfile.ZipFile(path) as z:
        inner = next(n for n in z.namelist() if n.lower().endswith((".xlsx", ".xls")))
        wb = openpyxl.load_workbook(io.BytesIO(z.read(inner)), read_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    hdr = [str(c) if c is not None else "" for c in next(it)]
    idx = {c.upper(): i for i, c in enumerate(hdr)}
    missing = [c for c in ("OCC_CODE", "O_GROUP", *WAGE_COLS.values())
               if c not in idx]
    if missing:
        raise SystemExit(f"  {path.name}: missing {missing}")
    out = {}
    for r in it:
        if str(r[idx["O_GROUP"]] or "").strip().lower() != DETAIL:
            continue
        code = str(r[idx["OCC_CODE"]] or "").strip()
        if not re.fullmatch(r"\d{2}-\d{4}", code):
            continue
        rec = {k: num(r[idx[col]]) for k, col in WAGE_COLS.items()}
        if rec["median"] is None:
            continue
        rec["title"] = str(r[idx.get("OCC_TITLE", 1)] or "").strip()
        emp = num(r[idx["TOT_EMP"]]) if "TOT_EMP" in idx else None
        rec["employment"] = int(emp) if emp else None
        out[code] = rec
    year = int(re.search(r"oesm(\d{2})nat", path.name).group(1)) + 2000
    return year, out


def main() -> int:
    files = sorted(SRC.glob("oesm*nat.zip"))
    if not files:
        print(f"  no OEWS files in {SRC.relative_to(ROOT)}")
        return 1

    years, warns = {}, []
    for p in files:
        y, data = read_year(p)
        if len(data) < 300:
            warns.append(f"{p.name}: only {len(data)} detailed occupations")
        years[y] = data
        print(f"  {y}: {len(data):,} detailed occupations")

    latest = max(years)
    cur = years[latest]

    # A sanity check with a known answer: the median must sit between the
    # 25th and 75th percentile in every occupation.
    for code, r in cur.items():
        if r["p25"] and r["p75"] and not (r["p25"] <= r["median"] <= r["p75"]):
            warns.append(f"{code}: median {r['median']} outside p25-p75")

    focus = []
    for code, label in FOCUS.items():
        r = cur.get(code)
        if not r:
            warns.append(f"{code} ({label}) not in {latest} data")
            continue
        focus.append({"soc": code, "occupation": r["title"] or label,
                      "employment": r["employment"], **{
                          k: r[k] for k in ("p10", "p25", "median", "p75", "p90")}})

    doc = {
        "_comment": (
            "Annual wage percentiles by detailed occupation, national, from "
            "the BLS Occupational Employment and Wage Statistics survey. THIS "
            "SAYS NOTHING ABOUT VISAS ON ITS OWN -- it is the wage "
            "distribution for the whole occupation, everyone in it, and "
            "exists here to give the H-1B wage-level figures a denominator in "
            "dollars. OEWS suppresses some cells ('*') and tops wages out at "
            "$239,200 ('#'); both are stored as null rather than guessed. "
            "Detailed six-digit SOC rows only, so the totals and major groups "
            "are excluded and the rows do not sum to national employment."),
        "source": "BLS Occupational Employment and Wage Statistics, national files",
        "source_url": "https://www.bls.gov/oes/tables.htm",
        "years": sorted(years),
        "latest_year": latest,
        "occupations": len(cur),
        "focus_occupations": focus,
        "by_soc": {y: {c: {k: r[k] for k in ("median", "p25", "p10", "p75", "p90")}
                       for c, r in d.items()}
                   for y, d in years.items()},
        "titles": {c: r["title"] for c, r in cur.items()},
        "warnings": warns,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(doc, indent=1) + "\n")

    print(f"\n  {latest} annual wages, occupations the H-1B programme runs through:")
    print(f"  {'occupation':44}{'p25':>10}{'median':>10}{'p75':>10}")
    for f in focus[:9]:
        p25 = f"${f['p25']:,.0f}" if f["p25"] else "-"
        med = f"${f['median']:,.0f}" if f["median"] else "-"
        p75 = f"${f['p75']:,.0f}" if f["p75"] else "-"
        print(f"  {f['occupation'][:42]:44}{p25:>10}{med:>10}{p75:>10}")
    for w in warns[:5]:
        print(f"  WARN {w}")
    print(f"\n  wrote {OUT.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
