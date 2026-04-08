#!/usr/bin/env python3
"""
Extract Canada study permit + IMP work permit data from IRCC Open Data
xlsx files into a single JSON for the canada.html page.

Output: data/canada/canada_dashboard.json
"""
import json
import warnings
import openpyxl
from pathlib import Path

warnings.filterwarnings("ignore")

ROOT = Path(__file__).resolve().parent.parent
CDN = ROOT / "data" / "canada"
OUT = CDN / "canada_dashboard.json"


def _parse_int(val):
    if val is None or val == "" or val == "--":
        return 0
    if isinstance(val, (int, float)):
        return int(val)
    s = str(val).replace(",", "").strip()
    try:
        return int(s)
    except ValueError:
        return 0


def _build_year_end_cols(rows):
    """Find the column index of the year-total cell for each year."""
    year_row = rows[2]
    last_year = None
    year_starts = {}
    year_ends = {}
    for i, c in enumerate(year_row):
        if c and isinstance(c, (int, float, str)):
            try:
                yr = int(str(c).strip())
                if 2000 <= yr <= 2030:
                    last_year = yr
            except ValueError:
                pass
        if last_year:
            if last_year not in year_starts:
                year_starts[last_year] = i
            year_ends[last_year] = i
    return year_ends


def load_study_permits_by_citizenship():
    """Yearly totals by country of citizenship."""
    wb = openpyxl.load_workbook(CDN / "study_permits_by_citizenship.xlsx", data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    year_ends = _build_year_end_cols(rows)
    years = sorted(year_ends.keys())

    # Skip 2026 (only Q1 partial data)
    years = [y for y in years if y < 2026]

    countries = {}
    for r in rows[5:]:
        if not r or not r[0] or not isinstance(r[0], str):
            continue
        country = r[0].strip()
        if country in ("Total", "Grand Total", "Country of Citizenship"):
            continue
        vals = []
        for y in years:
            col = year_ends[y]
            v = r[col] if col < len(r) else None
            vals.append(_parse_int(v))
        if any(vals):
            countries[country] = vals

    # Top countries by total over the period
    top = sorted(countries.items(), key=lambda kv: -sum(kv[1]))[:15]
    return {
        "years": years,
        "top_countries": [{"country": c, "yearly": v} for c, v in top],
        "all_countries_count": len(countries),
    }


def load_imp_yearend_by_citizenship():
    """IMP (work permit) holders on Dec 31 by country, 2000-2024."""
    wb = openpyxl.load_workbook(CDN / "imp_yearend_by_citizenship.xlsx", data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    # Header layout is different here — single header row 2 with year columns
    header = rows[2]
    year_cols = {}
    for i, c in enumerate(header):
        if c and isinstance(c, (int, float, str)):
            try:
                yr = int(str(c).strip())
                if 2000 <= yr <= 2030:
                    year_cols[yr] = i
            except ValueError:
                pass
    years = sorted(year_cols.keys())

    countries = {}
    for r in rows[3:]:
        if not r or not r[0] or not isinstance(r[0], str):
            continue
        country = r[0].strip()
        if country in ("Total", "Grand Total"):
            continue
        vals = [_parse_int(r[year_cols[y]]) if year_cols[y] < len(r) else 0 for y in years]
        if any(vals):
            countries[country] = vals

    top = sorted(countries.items(), key=lambda kv: -sum(kv[1]))[:15]
    return {
        "years": years,
        "top_countries": [{"country": c, "yearly": v} for c, v in top],
        "all_countries_count": len(countries),
    }


def load_study_permits_by_province_level():
    """Province × study level totals over the years."""
    wb = openpyxl.load_workbook(CDN / "study_permits_by_province_study_level.xlsx", data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    year_ends = _build_year_end_cols(rows)
    years = sorted(year_ends.keys())
    years = [y for y in years if y < 2026]

    # Province is in col 0 (sticky), study level is in col 1
    current_province = None
    out = {}  # province → {level: [yearly_totals]}
    for r in rows[5:]:
        if not r:
            continue
        prov_cell = r[0]
        if prov_cell and isinstance(prov_cell, str) and prov_cell.strip():
            current_province = prov_cell.strip()
        level = r[1] if r[1] else None
        if not level or not current_province:
            continue
        level = str(level).strip()
        if level in ("Province/Territory and study level",):
            continue
        vals = []
        for y in years:
            col = year_ends[y]
            vals.append(_parse_int(r[col]) if col < len(r) else 0)
        if any(vals):
            if current_province not in out:
                out[current_province] = {}
            out[current_province][level] = vals

    # Roll up: total post-secondary permits by province (latest year)
    latest_idx = -1
    province_postsec_latest = {}
    for prov, levels in out.items():
        if prov in ("Total", "Grand Total"):
            continue
        ps = levels.get("Post Secondary", [0] * len(years))
        province_postsec_latest[prov] = ps[latest_idx]
    top_provinces = sorted(province_postsec_latest.items(), key=lambda kv: -kv[1])[:10]

    # National totals by study level over time
    national = {}
    for prov, levels in out.items():
        if prov in ("Total", "Grand Total"):
            continue
        for level, vals in levels.items():
            if level not in national:
                national[level] = [0] * len(years)
            for i, v in enumerate(vals):
                national[level][i] += v

    return {
        "years": years,
        "top_provinces_postsec_latest": top_provinces,
        "national_by_level": national,
    }


def load_pgwp_by_province():
    """Post-graduation work permits by province over years."""
    wb = openpyxl.load_workbook(CDN / "pgwp_by_province_cma.xlsx", data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    year_ends = _build_year_end_cols(rows)
    years = sorted(year_ends.keys())
    years = [y for y in years if y < 2026]

    current_province = None
    province_totals = {}
    for r in rows[5:]:
        if not r:
            continue
        prov_cell = r[0]
        if prov_cell and isinstance(prov_cell, str) and prov_cell.strip():
            current_province = prov_cell.strip()
        cma = r[1] if r[1] else None
        if not cma or not current_province:
            continue
        if str(cma).strip() in ("Total", "Province/territory and CMA"):
            continue
        vals = []
        for y in years:
            col = year_ends[y]
            vals.append(_parse_int(r[col]) if col < len(r) else 0)
        if any(vals):
            if current_province not in province_totals:
                province_totals[current_province] = [0] * len(years)
            for i, v in enumerate(vals):
                province_totals[current_province][i] += v

    return {
        "years": years,
        "by_province": [
            {"province": p, "yearly": v}
            for p, v in sorted(province_totals.items(), key=lambda kv: -sum(kv[1]))
            if p not in ("Total", "Grand Total")
        ][:10],
    }


def main():
    out = {
        "source": "Statistics Canada / IRCC Open Data — study permit holders, IMP work permit holders, and Post-Graduation Work Permits",
        "url": "https://open.canada.ca/data/en/dataset/temporary-residents-by-immigration-category",
        "study_permits_by_citizenship": load_study_permits_by_citizenship(),
        "imp_yearend_by_citizenship": load_imp_yearend_by_citizenship(),
        "study_permits_by_province_level": load_study_permits_by_province_level(),
        "pgwp_by_province": load_pgwp_by_province(),
    }
    OUT.write_text(json.dumps(out, indent=2))
    print(f"Wrote {OUT}")
    print(f"  study permits years: {out['study_permits_by_citizenship']['years']}")
    print(f"  countries with data: {out['study_permits_by_citizenship']['all_countries_count']}")
    print(f"  IMP years: {out['imp_yearend_by_citizenship']['years'][:3]}...{out['imp_yearend_by_citizenship']['years'][-3:]}")


if __name__ == "__main__":
    main()
