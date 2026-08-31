#!/usr/bin/env python3
"""Visa issuances by class and nationality, 1997-2024.

Everything else in this project starts inside the United States: degrees
conferred, petitions filed, status adjusted. This is the entry point, counted
at the consulate, by nationality and by year. It is also the most current
series here -- State publishes through the last fiscal year, where IPEDS and
the SED run one to two years behind.

CLASS CODES ARE NOT SPELLED CONSISTENTLY across the twenty-eight sheets: the
early years write "F-1" and "H-1B", the later ones "F1" and "H1B". Codes are
normalised by stripping punctuation, or a series silently starts in 2010.

AN ISSUANCE IS NOT A PERSON AND NOT AN ARRIVAL. One person can be issued
several visas over the years, and a visa is permission to travel and seek
admission, not admission itself. These are counts of documents issued in a
fiscal year.

Usage:
  python3 scripts/build_visa_issuances.py
"""
from __future__ import annotations

import json
import re
import sys
import warnings
from collections import defaultdict
from pathlib import Path

warnings.filterwarnings("ignore")

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data" / "state_visas" / "1997-2024_NIVDetailTable.xlsx"
OUT = ROOT / "site" / "data" / "visa_issuances.json"

# The classes this project's argument runs through.
CLASSES = {"F1": "F-1 student", "F2": "F-2 student dependent",
           "J1": "J-1 exchange visitor", "H1B": "H-1B specialty occupation",
           "H4": "H-4 H-1B dependent", "L1": "L-1 intracompany transferee",
           "O1": "O-1 extraordinary ability", "TN": "TN professional",
           "E3": "E-3 Australian specialty", "M1": "M-1 vocational student"}
# Region headings and footers sit in the same column as country names.
NOT_A_COUNTRY = re.compile(
    r"^(africa|asia|europe|north america|south america|oceania|"
    r"grand total|totals?|unknown|non-?nationality|.*note.*)$", re.I)


def norm(code: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(code or "").upper())


# The State tables and the World Bank do not spell countries the same way.
COUNTRY_ALIASES = {
    "china - mainland": "china", "china - taiwan": "taiwan",
    "korea, south": "korea, rep.", "korea, north": "korea, dem. people's rep.",
    "russia": "russian federation", "vietnam": "viet nam",
    "egypt": "egypt, arab rep.", "iran": "iran, islamic rep.",
    "turkey": "turkiye", "venezuela": "venezuela, rb",
    "great britain and northern ireland": "united kingdom",
}


def norm_country(name: str) -> str:
    s = re.sub(r"\s+", " ", str(name or "")).strip().lower()
    return COUNTRY_ALIASES.get(s, s)


def load_population() -> dict[str, float]:
    """Country populations, plus 'world', from the cached World Bank pull."""
    path = ROOT / "data" / "worldbank" / "population_2024.json"
    if not path.exists():
        return {}
    payload = json.loads(path.read_text())
    if not isinstance(payload, list) or len(payload) < 2:
        return {}
    out: dict[str, float] = {}
    for r in payload[1]:
        name = (r.get("country") or {}).get("value")
        val = r.get("value")
        if name and isinstance(val, (int, float)):
            out[norm_country(name)] = float(val)
    if "world" in out:
        out["world"] = out["world"]
    return out



MONTHS = ("JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE", "JULY",
          "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER")


def parse_monthly() -> tuple[dict, list[str]]:
    """Monthly issuance files, which run ahead of the annual detail table.

    These are long format -- one row per nationality and class -- rather than
    the wide country-by-class grid of the annual sheets.
    """
    import openpyxl
    out: dict[str, dict] = {}
    warns: list[str] = []
    for path in sorted(SRC.parent.glob("* - NIV Issuances by Nationality*.xlsx")):
        name = path.name.upper()
        month = next((m for m in MONTHS if m in name), None)
        ym = re.search(r"(20\d\d)", name)
        if not month or not ym:
            warns.append(f"cannot date {path.name}")
            continue
        key = f"{ym.group(1)}-{MONTHS.index(month) + 1:02d}"
        ws = openpyxl.load_workbook(path, read_only=True).active
        rows = ws.iter_rows(values_only=True)
        hdr = None
        by_class: dict[str, int] = defaultdict(int)
        for r in rows:
            cells = [str(x).strip() if x is not None else "" for x in r[:3]]
            if hdr is None:
                if cells[:2] == ["Nationality", "Visa Class"]:
                    hdr = True
                continue
            code = norm(cells[1])
            if code in CLASSES and isinstance(r[2], (int, float)):
                by_class[code] += int(r[2])
        if by_class:
            out[key] = dict(sorted(by_class.items()))
        else:
            warns.append(f"{path.name}: no rows matched")
    return dict(sorted(out.items())), warns


def main() -> int:
    if not SRC.exists():
        print(f"  missing {SRC.relative_to(ROOT)}")
        return 1
    import openpyxl
    wb = openpyxl.load_workbook(SRC, read_only=True)

    totals = defaultdict(lambda: defaultdict(int))        # year -> class -> n
    by_country = defaultdict(lambda: defaultdict(dict))   # class -> country -> year -> n
    warns = []

    for sheet in wb.sheetnames:
        m = re.search(r"(\d{2,4})", sheet)
        if not m:
            continue
        yy = int(m.group(1))
        year = yy if yy > 1900 else (1900 + yy if yy >= 90 else 2000 + yy)
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        hdr = [norm(x) for x in next(rows)]
        cols = {c: i for i, c in enumerate(hdr) if c in CLASSES}
        missing = set(CLASSES) - set(cols)
        if missing:
            warns.append(f"{year}: no column for {sorted(missing)}")
        for r in rows:
            name = str(r[0] or "").strip()
            if not name or NOT_A_COUNTRY.match(name):
                continue
            for code, i in cols.items():
                v = r[i] if i < len(r) else None
                if isinstance(v, (int, float)):
                    totals[year][code] += int(v)
                    by_country[code][name][year] = int(v)

    years = sorted(totals)
    if len(years) < 10:
        print(f"  only {len(years)} years parsed")
        return 1
    f, l = years[0], years[-1]

    # Leading origin countries, ranked on the latest year of each class.
    top = {}
    for code in ("F1", "H1B", "J1"):
        rank = sorted(by_country[code].items(),
                      key=lambda kv: -kv[1].get(l, 0))[:12]
        tot_l = totals[l][code]
        top[CLASSES[code]] = [
            {"country": name, "issuances": ys.get(l, 0),
             "share_pct": round(100 * ys.get(l, 0) / tot_l, 1) if tot_l else None,
             "change_since_first_pct": (
                 round(100 * (ys.get(l, 0) / ys[f] - 1), 1)
                 if ys.get(f) else None)}
            for name, ys in rank]

    # The monthly files carry the current fiscal year, which the annual
    # detail table does not reach yet. Long format: nationality, class, count.
    monthly, m_warns = parse_monthly()

    # Overrepresentation needs a denominator. Population is the plainest one
    # a reader can check, and it is not the only defensible choice -- share
    # of the world's STEM graduates would be another -- so the ratio is
    # labelled as being against population and nothing else.
    pop = load_population()
    over = {}
    for code in ("F1", "H1B"):
        tot = totals[l].get(code, 0)
        if not tot or not pop.get("world"):
            continue
        rows = []
        for r in top[CLASSES[code]][:8]:
            wb = pop.get(norm_country(r["country"]))
            if not wb:
                continue
            pop_share = 100 * wb / pop["world"]
            rows.append({"country": r["country"],
                         "issuance_share_pct": r["share_pct"],
                         "population_share_pct": round(pop_share, 2),
                         "ratio": round(r["share_pct"] / pop_share, 1)})
        over[CLASSES[code]] = rows

    doc = {
        "_comment": (
            "Nonimmigrant visa issuances by class and nationality, fiscal "
            "years 1997-2024, from the State Department detail tables. AN "
            "ISSUANCE IS NOT A PERSON AND NOT AN ARRIVAL: one person may be "
            "issued several visas over the years, and a visa is permission to "
            "travel and seek admission, not admission itself. Nationality is "
            "the applicant's, not the consulate's country. Class codes are "
            "written with hyphens in the early sheets and without in the "
            "later ones, and are normalised here. This is the most current "
            "series in the project -- State publishes through the last fiscal "
            "year while IPEDS and the SED run one to two years behind."),
        "source": "US Department of State, Nonimmigrant Visa Detail Tables, FY1997-2024",
        "source_url": "https://travel.state.gov/content/travel/en/legal/visa-law0/visa-statistics/nonimmigrant-visa-statistics.html",
        "classes": CLASSES,
        "years": years,
        "totals": {str(y): dict(totals[y]) for y in years},
        "leading_origins": top,
        "monthly": monthly,
        "overrepresentation_vs_population": over,
        "population_source": "World Bank, SP.POP.TOTL, 2024",
        "warnings": warns + m_warns,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(doc, indent=1) + "\n")

    print(f"  {len(years)} years, {f}-{l}, {len(by_country['F1'])} countries")
    print(f"  {'year':6}" + "".join(f"{CLASSES[c].split()[0]:>10}" for c in
                                    ("F1", "J1", "H1B", "L1", "O1")))
    for y in years:
        if y % 5 == 0 or y in (f, l):
            print(f"  {y:6}" + "".join(f"{totals[y].get(c, 0):>10,}"
                                       for c in ("F1", "J1", "H1B", "L1", "O1")))
    print(f"\n  leading F-1 origins in {l}:")
    for r in top["F-1 student"][:5]:
        print(f"   {r['country']:22}{r['issuances']:>9,}  {r['share_pct']:>5}%")
    print(f"  leading H-1B origins in {l}:")
    for r in top["H-1B specialty occupation"][:5]:
        print(f"   {r['country']:22}{r['issuances']:>9,}  {r['share_pct']:>5}%")
    if over:
        for cls, rows in over.items():
            print(f"\n  {cls}: share of issuances against share of world population")
            for r in rows[:4]:
                print(f"   {r['country']:22}{r['issuance_share_pct']:>6}% vs "
                      f"{r['population_share_pct']:>5}%  = {r['ratio']:>5}x")
    if monthly:
        print(f"\n  monthly, {min(monthly)} to {max(monthly)}:")
        print(f"  {'month':10}" + "".join(f"{c:>9}" for c in ("F1", "J1", "H1B")))
        for k, v in monthly.items():
            print(f"  {k:10}" + "".join(f"{v.get(c, 0):>9,}" for c in ("F1", "J1", "H1B")))
    for w in (warns + m_warns)[:6]:
        print(f"  WARN {w}")
    print(f"  wrote {OUT.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
