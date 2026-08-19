#!/usr/bin/env python3
"""F-1 student visa issuance by fiscal year, from the State Department NIV table.

Source: data/h1b/state_dept/FYs97-24_NIVDetailTable.xlsx -- one worksheet per
fiscal year, countries as rows, visa classes as columns.

TWO TRAPS IN THIS FILE, both of which produce wrong answers silently:

1. Each sheet mixes three kinds of row: individual countries, regional
   subtotals ("Totals for Asia"), and in most years a "Grand Totals" row.
   Summing every row triple-counts. A naive sum gives 1,932,699 for FY2015
   against a true figure of 644,233 -- almost exactly 3x, which is just close
   enough to a plausible number to go unnoticed.

2. The column naming changes partway through. FY1997-FY2021 use hyphenated
   codes ('F-1'); FY2022 onward use unhyphenated ('F1'). A reader matching
   only 'F-1' silently loses the three most recent years -- the ones most
   likely to matter. Those later sheets also drop the "Grand Totals" row, so
   the total has to be built from the regional subtotals instead.

Verification anchor: FY2015 F-1 issuance is 644,233. If a change to this
script moves that number, the change is wrong.

Usage:
    python3 scripts/build_f1_series.py                 # print the series
    python3 scripts/build_f1_series.py --json PATH     # write JSON
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "h1b" / "state_dept" / "FYs97-24_NIVDetailTable.xlsx"

ANCHOR_YEAR, ANCHOR_VALUE = 2015, 644_233

# Countries broken out individually in the output.
WATCH = {"china - mainland": "china", "india": "india",
         "korea, south": "south_korea", "saudi arabia": "saudi_arabia",
         "nigeria": "nigeria", "vietnam": "vietnam"}


def fiscal_year(sheet_name: str) -> int | None:
    m = re.match(r"FY(\d{2})$", sheet_name)
    if not m:
        return None
    yy = int(m.group(1))
    return 1900 + yy if yy >= 97 else 2000 + yy


def f1_column(header) -> int | None:
    """The F-1 column, under either naming convention."""
    for want in ("F-1", "F1"):
        for i, cell in enumerate(header):
            if cell is not None and str(cell).strip() == want:
                return i
    return None


def parse_sheet(ws):
    """Return (total, {country_key: n}) for one fiscal year."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return None, {}
    col = f1_column(rows[0])
    if col is None:
        return None, {}

    grand = None
    region_sum = 0
    watched: dict[str, int] = {}
    for row in rows[1:]:
        name = str(row[0]).strip() if row and row[0] else ""
        if not name or len(row) <= col or row[col] is None:
            continue
        try:
            n = int(str(row[col]).replace(",", "").strip())
        except ValueError:
            continue
        low = name.lower()
        if low.startswith("grand total"):
            grand = n
        elif low.startswith("totals for"):
            # Regional subtotal: the fallback when no grand total exists.
            region_sum += n
        elif low in WATCH:
            watched[WATCH[low]] = n
    return (grand if grand is not None else (region_sum or None)), watched


def build():
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl required: pip install openpyxl")
    if not SRC.exists():
        sys.exit(f"Source not found: {SRC}")
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

    out: dict[int, dict] = {}
    for sheet in wb.sheetnames:
        year = fiscal_year(sheet)
        if year is None:
            continue
        total, watched = parse_sheet(wb[sheet])
        if total is None:
            continue
        entry = {"total": total}
        entry.update(watched)
        out[year] = entry
    return dict(sorted(out.items()))


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--json", help="write the series to this path")
    args = ap.parse_args()

    series = build()
    if not series:
        sys.exit("No F-1 data extracted.")

    got = series.get(ANCHOR_YEAR, {}).get("total")
    if got != ANCHOR_VALUE:
        print(f"WARNING: FY{ANCHOR_YEAR} F-1 is {got:,} but the verified "
              f"figure is {ANCHOR_VALUE:,}. The parse is wrong.", file=sys.stderr)

    print(f"{'FY':<6}{'F-1 issued':>12}{'China':>10}{'India':>10}{'CN+IN':>8}")
    print("-" * 46)
    for year, e in series.items():
        t = e["total"]
        c, i = e.get("china", 0), e.get("india", 0)
        share = f"{100 * (c + i) / t:.0f}%" if t else "-"
        print(f"{year:<6}{t:>12,}{c:>10,}{i:>10,}{share:>8}")

    years = list(series)
    peak = max(series, key=lambda y: series[y]["total"])
    print(f"\npeak {series[peak]['total']:,} in FY{peak}; "
          f"latest FY{years[-1]} {series[years[-1]]['total']:,} "
          f"({100 * (series[years[-1]]['total'] - series[peak]['total']) / series[peak]['total']:+.0f}% off peak)")

    if args.json:
        Path(args.json).write_text(json.dumps(series, indent=2))
        print(f"\nwrote {args.json}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
