#!/usr/bin/env python3
"""L-1, O-1 and the rest: admissions by class, year and destination state.

There is no employer-level data for L-1 or O-1 -- neither requires a labour
condition application and USCIS publishes no employer hub for either -- so
the question "which companies bring in the most" has no published answer.
This is the closest available substitute: how many were admitted, in which
years, and to which states.

ADMISSIONS ARE NOT PEOPLE AND NOT ISSUANCES. An admission is a single entry
at a port, so one worker who travels four times in a year is four
admissions. That makes this series much larger than the visa-issuance series
already held, and the two must never be mixed: issuances count documents,
admissions count arrivals, and neither counts workers.

STATE IS THE DESTINATION GIVEN ON ARRIVAL, which for an intracompany
transferee is usually the office they are joining. It is the nearest thing to
an employer signal available for L-1 and O-1, and it is not an employer.

Usage:
  python3 scripts/build_nonimmigrant_admissions.py
"""
from __future__ import annotations

import json
import re
import sys
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data" / "dhs_yearbook" / "nonimmigrants_fy2023.xlsx"
OUT = ROOT / "site" / "data" / "nonimmigrant_admissions.json"
sys.path.insert(0, str(ROOT / "scripts"))
from atomic_write import write_json                          # noqa: E402

# The classes this project follows. Dependants are carried because an L-2 or
# H-4 admission is a family arriving on the back of a worker's status.
CLASSES = {
    "H1B": "H-1B specialty occupation", "H4": "H-4 dependant",
    "L1": "L-1 intracompany transferee", "L2": "L-2 dependant",
    "O1": "O-1 extraordinary ability", "O2": "O-2 support",
    "O3": "O-3 dependant", "F1": "F-1 student", "F2": "F-2 dependant",
    "J1": "J-1 exchange visitor", "TN": "TN professional",
    "E3": "E-3 Australian specialty",
}
NOT_A_STATE = {"total", "class", "description", "unknown", "other"}


def norm(code) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(code or "").upper())


def main() -> int:
    if not SRC.exists():
        print(f"  missing {SRC.relative_to(ROOT)}")
        return 1
    import openpyxl
    wb = openpyxl.load_workbook(SRC, read_only=True)

    # --- Table 25: admissions by class and year -------------------------
    ws = wb["Table 25"]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    years, hdr_at = [], None
    for i, r in enumerate(rows[:12]):
        cand = [str(x).strip() for x in r[1:] if x is not None]
        if len(cand) >= 5 and all(re.fullmatch(r"(19|20)\d\d", c) for c in cand):
            years, hdr_at = cand, i
            break
    if not years:
        print("  no year header in Table 25")
        return 1

    # The class code sits inside the label, e.g. "Intracompany transferees
    # (L1)". Matching the parenthetical is what identifies the row.
    by_class: dict[str, dict[str, int]] = {}
    for r in rows[hdr_at + 1:]:
        label = re.sub(r"\s+", " ", str(r[0] or "")).strip()
        m = re.search(r"\(([A-Z][A-Z0-9, and/-]*)\)\s*$", label)
        if not m:
            continue
        codes = [norm(c) for c in re.split(r"[,/]| and ", m.group(1))]
        code = next((c for c in codes if c in CLASSES), None)
        if not code:
            continue
        vals = {y: int(v) for y, v in zip(years, r[1:])
                if isinstance(v, (int, float))}
        if vals:
            by_class.setdefault(code, {}).update(vals)

    # --- NISuppTable3: admissions by class and destination state --------
    ws3 = wb["NISuppTable3"]
    rows3 = [list(r) for r in ws3.iter_rows(values_only=True)]
    hdr3 = next((r for r in rows3[:10]
                 if str(r[0] or "").strip().lower() == "class"), None)
    by_state: dict[str, dict[str, int]] = {}
    if hdr3:
        states = [(i, str(v).strip()) for i, v in enumerate(hdr3)
                  if v and str(v).strip().lower() not in NOT_A_STATE]
        for r in rows3:
            code = norm(r[0])
            if code not in CLASSES:
                continue
            got = {name: int(r[i]) for i, name in states
                   if i < len(r) and isinstance(r[i], (int, float))}
            if got:
                by_state[code] = dict(sorted(got.items(), key=lambda kv: -kv[1]))

    if not by_class:
        print("  no classes matched in Table 25")
        return 1

    warns = []
    for code, ys in by_class.items():
        st = by_state.get(code)
        if not st:
            continue
        latest = max(ys)
        # The state table is a single year; it should not exceed that year's
        # national admissions for the same class by more than rounding.
        tot = sum(v for k, v in st.items() if k.lower() != "total")
        if ys.get(latest) and tot > ys[latest] * 1.05:
            warns.append(f"{code}: state sum {tot:,} exceeds {latest} "
                         f"national {ys[latest]:,}")

    doc = {
        "_comment": (
            "Nonimmigrant ADMISSIONS by class of admission and year (DHS "
            "Yearbook table 25), with admissions by class and destination "
            "state (supplemental table 3). Built because there is NO "
            "EMPLOYER-LEVEL DATA FOR L-1 OR O-1: neither requires a labour "
            "condition application and USCIS publishes no employer hub for "
            "either, so this is the closest available substitute and it is "
            "not a substitute for an employer. AN ADMISSION IS AN ENTRY AT A "
            "PORT, not a person and not a visa: one worker travelling four "
            "times is four admissions, which is why these numbers far exceed "
            "the issuance series held separately. NEVER MIX THE TWO -- "
            "issuances count documents, admissions count arrivals, neither "
            "counts workers. STATE IS THE DESTINATION DECLARED ON ARRIVAL, "
            "usually the office being joined for a transferee; it is the "
            "nearest thing to an employer signal for these classes and it is "
            "not an employer. Figures are rounded by DHS."),
        "source": "DHS Office of Homeland Security Statistics, Yearbook of "
                  "Immigration Statistics, nonimmigrant tables",
        "source_url": "https://ohss.dhs.gov/topics/immigration/yearbook",
        "classes": {c: CLASSES[c] for c in sorted(by_class)},
        "years": [int(y) for y in years],
        "by_class": {c: by_class[c] for c in sorted(by_class)},
        "by_destination_state": {c: dict(list(by_state[c].items())[:15])
                                 for c in sorted(by_state)},
        "not_published_by_employer": ["L-1", "O-1"],
        "warnings": warns,
    }
    write_json(OUT, doc)

    f, l = years[0], years[-1]
    print(f"  {len(by_class)} classes, {f}-{l}")
    print(f"  {'class':34}{f:>12}{l:>12}{'change':>9}")
    for code in ("H1B", "L1", "O1", "F1", "J1", "TN"):
        ys = by_class.get(code)
        if not ys or f not in ys or l not in ys:
            continue
        ch = 100 * (ys[l] / ys[f] - 1) if ys[f] else 0
        print(f"  {CLASSES[code][:32]:34}{ys[f]:>12,}{ys[l]:>12,}{ch:>8.0f}%")
    for code in ("L1", "O1"):
        st = by_state.get(code)
        if st:
            top = [f"{k} {v:,}" for k, v in list(st.items())[:5]
                   if k.lower() != "total"]
            print(f"\n  {CLASSES[code]} by destination state: {', '.join(top)}")
    for w in warns[:5]:
        print(f"  WARN {w}")
    print(f"\n  wrote {OUT.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
